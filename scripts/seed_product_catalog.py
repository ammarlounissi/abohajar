"""Import commercial names, sizes, descriptions and search tags from the product XLSX.

Run from the project root after applying migrations:
    python -m scripts.seed_product_catalog
    python -m scripts.seed_product_catalog --dry-run
    python -m scripts.seed_product_catalog --xlsx /path/to/file.xlsx

Only existing products are updated (matched by SKU); this script never inserts products,
changes category_code, prices, media, or factory ownership. Blank spreadsheet fields are
non-destructive. Use --rewrite-tag-encoding once to rewrite legacy escaped JSON text.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from sqlalchemy.orm.attributes import flag_modified

from app import models
from app.core.database import SessionLocal

DEFAULT_XLSX = Path(__file__).resolve().parents[1] / "product_catalog_seed.xlsx"
SHEET_NAME = "المنتجات"


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\u00a0", " ").strip().split())


def _header_key(value: Any) -> str:
    return _clean(value).casefold().replace(" ", "")


def _first(row: dict[str, Any], *keys: str) -> str:
    normalized_row = {_header_key(key): value for key, value in row.items()}
    for key in keys:
        value = _clean(normalized_row.get(_header_key(key)))
        if value:
            return value
    return ""


def _split_multi(value: str) -> list[str]:
    """Split bilingual spreadsheet cells on comma-like separators, preserving words."""
    normalized = value.replace("،", ",").replace(";", ",").replace("|", ",")
    return [part.strip() for part in normalized.split(",") if part.strip()]


def build_tags(row: dict[str, Any]) -> list[str]:
    """Derive compact, factual tags from the structured columns in the XLSX row."""
    category = _first(row, "الصنف", "الفئة")
    stone_type = _first(row, "نوع الحجر الطبيعي", "نوع الحجر")
    color = _first(row, "اللون الأساسي", "اللون")
    texture = _first(row, "الملمس")
    pattern = _first(row, "نمط اللون")
    size = _first(row, "المقاس", "المقاسات")
    usage = _first(row, "الاستخدام")
    surface = _first(row, "السطح", "نوع السطح")
    title = _first(row, "العنوان الأصلي", "العنوان", "title")

    tags: list[str] = []

    def add(tag: str) -> None:
        tag = _clean(tag)
        if tag and len(tag) <= 50 and tag not in tags and len(tags) < 30:
            tags.append(tag)

    normalized_category = category.replace(" ", "")
    if normalized_category in {"طبيعي", "طبيعية", "حجرطبيعي"}:
        add("حجر طبيعي")
    elif normalized_category in {"صناعي", "صناعية", "حجرصناعي"}:
        add("حجر صناعي")
    elif category:
        # For accessories/materials (e.g. varnish), preserve their actual class without
        # misleadingly tagging them as a stone veneer.
        add(category)

    if stone_type and normalized_category not in {"صناعي", "صناعية", "حجرصناعي"}:
        add(stone_type)

    if color:
        if pattern == "مختلط":
            add("ألوان مختلطة")
            add(f"درجات {color}")
        else:
            add(color)

    if texture:
        add(f"ملمس {texture}")

    if size and size.casefold() not in {"عشوائي", "غير محدد", "-", "—"}:
        # Weight is not a dimension; keep the V001 2.5kg value out of size tags.
        if not any(unit in size.casefold() for unit in ("كغ", "kg", "g")):
            add(f"مقاس {size}")

    for item in _split_multi(usage):
        add(item)
    surface_names = {"ارضيات": "أرضيات", "جدران": "جدران"}
    for item in _split_multi(surface):
        add(surface_names.get(item, item))

    # A useful fallback for products whose structured fields are mostly unfilled.
    if not tags and title:
        add(title)
    return tags


def read_catalog(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"ملف XLSX غير موجود: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = next(values, None)
        if not headers:
            raise ValueError("ورقة المنتجات فارغة")
        keyed_headers = [_header_key(value) for value in headers]
        if not any(key in keyed_headers for key in ("sku", "رمزالمنتج")):
            raise ValueError("الملف يجب أن يحتوي عمود SKU")

        rows: list[dict[str, Any]] = []
        seen: set[str] = set()
        for excel_row, values_row in enumerate(values, start=2):
            cells = dict(zip(keyed_headers, values_row))
            row = {_header_key(k): v for k, v in cells.items()}
            sku = _first(row, "SKU", "رمز المنتج").upper()
            if not sku:
                continue
            if sku in seen:
                raise ValueError(f"رمز SKU مكرر في الملف: {sku} (حتى السطر {excel_row})")
            seen.add(sku)
            commercial_name = _first(row, "الاسم التجاري", "commercial_name")
            sizes = _first(row, "المقاس", "المقاسات", "sizes")
            description = _first(row, "الوصف", "description")
            # A package weight (e.g. kg) is not a stone size/dimension.
            if any(unit in sizes.casefold() for unit in ("كغ", "kg", "g")):
                sizes = ""
            raw_tags = _first(row, "الوسوم", "tags")
            tags = _split_multi(raw_tags) if raw_tags else build_tags(row)
            # De-duplicate and enforce model-side bounds.
            tags = list(dict.fromkeys(tag for tag in (_clean(t) for t in tags) if tag))[:30]
            too_long = [tag for tag in tags if len(tag) > 50]
            if too_long:
                raise ValueError(f"وسم أطول من 50 حرفاً لـ {sku}: {too_long[0]}")
            rows.append({
                "sku": sku,
                "commercial_name": commercial_name,
                "sizes": sizes,
                "description": description,
                "tags": tags,
            })
        return rows
    finally:
        workbook.close()


def seed_catalog(
    db: Session,
    rows: list[dict[str, Any]],
    *,
    dry_run: bool = False,
    rewrite_tag_encoding: bool = False,
) -> dict[str, Any]:
    summary = {"matched": 0, "updated": 0, "unchanged": 0, "tags_reencoded": 0, "missing": []}
    for item in rows:
        product = db.scalar(select(models.Product).where(models.Product.sku == item["sku"]))
        if product is None:
            summary["missing"].append(item["sku"])
            continue
        summary["matched"] += 1
        changed = False
        # Blank XLSX fields are intentionally non-destructive.
        for field in ("commercial_name", "sizes", "description"):
            value = item[field]
            if value and getattr(product, field) != value:
                setattr(product, field, value)
                changed = True
        if item["tags"] and product.tags != item["tags"]:
            product.tags = item["tags"]
            changed = True
        elif item["tags"] and rewrite_tag_encoding:
            # Force one UPDATE so SQLite stores Arabic JSON characters literally with
            # the project's ensure_ascii=False serializer; the value remains unchanged.
            flag_modified(product, "tags")
            summary["tags_reencoded"] += 1
        if changed:
            product.sync_status = models.SyncStatus.PENDING
            summary["updated"] += 1
        else:
            summary["unchanged"] += 1
    if dry_run:
        db.rollback()
    else:
        db.commit()
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description="استيراد الاسم التجاري والمقاسات والوصف والوسوم من ملف XLSX")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX, help="مسار ملف XLSX")
    parser.add_argument("--dry-run", action="store_true", help="محاكاة دون حفظ التغييرات")
    parser.add_argument(
        "--rewrite-tag-encoding",
        action="store_true",
        help="إعادة حفظ الوسوم المطابقة لعرض العربية مباشرة في تخزين JSON (نفّذها مرة واحدة)",
    )
    args = parser.parse_args()

    try:
        rows = read_catalog(args.xlsx)
    except (FileNotFoundError, ValueError) as exc:
        print(f"خطأ: {exc}", file=sys.stderr)
        return 2

    db = SessionLocal()
    try:
        summary = seed_catalog(
            db, rows, dry_run=args.dry_run, rewrite_tag_encoding=args.rewrite_tag_encoding,
        )
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()

    mode = "محاكاة فقط" if args.dry_run else "تم الحفظ"
    print(
        f"{mode}: {len(rows)} صفاً؛ مطابق {summary['matched']}؛ معدّل {summary['updated']}؛ "
        f"دون تغيير {summary['unchanged']}؛ أعيد ترميز وسوم {summary['tags_reencoded']}."
    )
    if summary["missing"]:
        print(f"SKU غير موجود في قاعدة البيانات ({len(summary['missing'])}): {', '.join(summary['missing'][:50])}")
        if len(summary["missing"]) > 50:
            print("... توجد رموز إضافية غير معروضة")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
